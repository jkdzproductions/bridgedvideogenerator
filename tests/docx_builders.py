"""Build small synthetic .docx files for the script_input tests.

python-docx has no API for hyperlinks, so the two link helpers write the WordprocessingML by
hand: `add_hyperlink` makes the `w:hyperlink` element Word writes for Insert > Link, and
`add_field_hyperlink` makes the HYPERLINK field code Word writes when a link is pasted from a
browser or another app (w:fldChar begin / w:instrText / separate / result runs / end).
"""
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def new_document():
    return Document()


def add_run(paragraph, text, bold=False, italic=False, style=None):
    run = paragraph.add_run(text)
    if bold:
        run.bold = True
    if italic:
        run.italic = True
    if style is not None:
        run.style = style
    return run


def add_hyperlink(paragraph, url, parts):
    """parts: list of (text, italic) or (text, italic, bold) runs that all sit inside one w:hyperlink."""
    rel_id = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    for part in parts:
        text, italic = part[0], part[1]
        bold = part[2] if len(part) > 2 else False
        run = paragraph.add_run(text)  # created at the end of the paragraph...
        run.italic = italic or None
        run.bold = bold or None
        hyperlink.append(run._r)  # ...then moved into the hyperlink element
    paragraph._p.append(hyperlink)


def _field_char(paragraph, kind):
    run = paragraph.add_run()
    fld = OxmlElement("w:fldChar")
    fld.set(qn("w:fldCharType"), kind)
    run._r.append(fld)


def add_field_hyperlink(paragraph, url, parts):
    """The same link written as a HYPERLINK field code instead of a w:hyperlink element.

    parts: list of (text, italic) or (text, italic, bold) runs."""
    _field_char(paragraph, "begin")
    instr_run = paragraph.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f' HYPERLINK "{url}" '
    instr_run._r.append(instr)
    _field_char(paragraph, "separate")
    for part in parts:
        text, italic = part[0], part[1]
        bold = part[2] if len(part) > 2 else False
        run = paragraph.add_run(text)
        run.italic = italic or None
        run.bold = bold or None
    _field_char(paragraph, "end")


def bold_character_style(document, name="Graph Bold"):
    """A character style whose font is bold — bold applied through a style, not the B button."""
    style = document.styles.add_style(name, WD_STYLE_TYPE.CHARACTER)
    style.font.bold = True
    return style


def italic_character_style(document, name="Graph Italic"):
    """A character style whose font is italic — italic applied through a style, not the I button."""
    style = document.styles.add_style(name, WD_STYLE_TYPE.CHARACTER)
    style.font.italic = True
    return style


def save(document, tmp_path, name="script.docx"):
    path = tmp_path / name
    document.save(str(path))
    return str(path)
